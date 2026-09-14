"""附件下载与文本解析（T3）。

支持 pdf（PyMuPDF）/ docx（python-docx）/ xlsx（openpyxl）；
旧版 .doc/.xls/.ppt 无解析库，标记 skipped；解析库懒加载导入，
缺失时对应附件标记 failed 而不中断任务。
"""
from __future__ import annotations

import uuid
from pathlib import Path
from urllib.parse import urlparse

import httpx

from ..config import ATTACHMENT_DIR, load_settings

ATTACHMENT_EXTS = {"pdf", "doc", "docx", "xls", "xlsx", "ppt", "pptx"}
PARSEABLE_EXTS = {"pdf", "docx", "xlsx"}
MAX_TEXT_LEN = 200_000


def attachment_ext(url: str) -> str:
    suffix = Path(urlparse(url).path).suffix.lower().lstrip(".")
    return suffix if suffix in ATTACHMENT_EXTS else ""


async def download(url: str) -> tuple[Path, str, int]:
    """下载到 data/attachments，返回 (本地路径, 扩展名, 字节数)。"""
    settings = load_settings()
    ext = attachment_ext(url) or "bin"
    ATTACHMENT_DIR.mkdir(parents=True, exist_ok=True)
    dest = ATTACHMENT_DIR / f"{uuid.uuid4().hex}.{ext}"
    timeout = float(settings.get("http_timeout", 20))
    total = 0
    async with httpx.AsyncClient(
        timeout=timeout, follow_redirects=True, headers={"User-Agent": settings.get("user_agent")}
    ) as client:
        async with client.stream("GET", url) as resp:
            resp.raise_for_status()
            with dest.open("wb") as fh:
                async for chunk in resp.aiter_bytes():
                    fh.write(chunk)
                    total += len(chunk)
    return dest, ext, total


def parse_text(path: Path, ext: str) -> tuple[str, str]:
    """返回 (解析状态, 文本)；状态为 parsed/failed/skipped。"""
    ext = ext.lower()
    if ext not in PARSEABLE_EXTS:
        return "skipped", ""
    try:
        if ext == "pdf":
            return "parsed", _parse_pdf(path)
        if ext == "docx":
            return "parsed", _parse_docx(path)
        if ext == "xlsx":
            return "parsed", _parse_xlsx(path)
    except Exception as exc:  # 解析失败不致命
        return "failed", f"{type(exc).__name__}: {exc}"
    return "skipped", ""


def _parse_pdf(path: Path) -> str:
    import fitz  # PyMuPDF

    parts: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
            if sum(len(p) for p in parts) > MAX_TEXT_LEN:
                break
    return "\n".join(parts)[:MAX_TEXT_LEN]


def _parse_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)[:MAX_TEXT_LEN]


def _parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
            if sum(len(p) for p in parts) > MAX_TEXT_LEN:
                break
    wb.close()
    return "\n".join(parts)[:MAX_TEXT_LEN]
