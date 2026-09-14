"""统一结果列表：筛选、分页、清空、Excel 导出。"""
from __future__ import annotations

import io
from datetime import datetime

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..database import get_db
from ..models import Platform, SearchResult
from .. import schemas

router = APIRouter(prefix="/results", tags=["results"])


def _apply_filters(stmt, *, keyword, platform_id, title, link, task_id, is_external):
    if keyword:
        stmt = stmt.where(SearchResult.keyword.contains(keyword))
    if platform_id is not None:
        stmt = stmt.where(SearchResult.platform_id == platform_id)
    if title:
        stmt = stmt.where(SearchResult.title.contains(title))
    if link:
        stmt = stmt.where(SearchResult.url.contains(link))
    if task_id is not None:
        stmt = stmt.where(SearchResult.task_id == task_id)
    if is_external is not None:
        stmt = stmt.where(SearchResult.is_external.is_(is_external))
    return stmt


def _serialize(row: tuple[SearchResult, str]) -> dict:
    result, platform_name = row
    return {
        "id": result.id,
        "task_id": result.task_id,
        "platform_id": result.platform_id,
        "keyword": result.keyword,
        "page": result.page,
        "title": result.title,
        "url": result.url,
        "is_external": result.is_external,
        "content_text": result.content_text,
        "screenshot_path": result.screenshot_path,
        "searched_at": result.searched_at,
        "platform_name": platform_name,
    }


@router.get("")
def list_results(
    keyword: str | None = None,
    platform_id: int | None = None,
    title: str | None = None,
    link: str | None = None,
    task_id: int | None = None,
    is_external: bool | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=200),
    db: Session = Depends(get_db),
):
    stmt = select(SearchResult, Platform.name).join(
        Platform, Platform.id == SearchResult.platform_id, isouter=True
    ).order_by(SearchResult.id.desc())
    stmt = _apply_filters(
        stmt,
        keyword=keyword,
        platform_id=platform_id,
        title=title,
        link=link,
        task_id=task_id,
        is_external=is_external,
    )
    total = len(db.execute(stmt).all())
    rows = db.execute(
        stmt.limit(page_size).offset((page - 1) * page_size)
    ).all()
    return {"total": total, "items": [_serialize(r) for r in rows]}


@router.delete("")
def clear_results(
    keyword: str | None = None,
    platform_id: int | None = None,
    title: str | None = None,
    link: str | None = None,
    task_id: int | None = None,
    is_external: bool | None = None,
    db: Session = Depends(get_db),
):
    """一键清空（带确认在前端）；带筛选参数时只清空当前筛选结果。"""
    stmt = select(SearchResult)
    stmt = _apply_filters(
        stmt,
        keyword=keyword,
        platform_id=platform_id,
        title=title,
        link=link,
        task_id=task_id,
        is_external=is_external,
    )
    rows = list(db.execute(stmt).scalars())
    for row in rows:
        db.delete(row)
    db.commit()
    return {"message": f"已清空 {len(rows)} 条结果", "deleted": len(rows)}


@router.get("/export")
def export_results(
    keyword: str | None = None,
    platform_id: int | None = None,
    title: str | None = None,
    link: str | None = None,
    task_id: int | None = None,
    is_external: bool | None = None,
    db: Session = Depends(get_db),
):
    """导出当前筛选结果为 Excel（含敏感命中分级列）。"""
    from ..models import SensitiveHit

    stmt = select(SearchResult, Platform.name).join(
        Platform, Platform.id == SearchResult.platform_id, isouter=True
    ).order_by(SearchResult.id.desc())
    stmt = _apply_filters(
        stmt,
        keyword=keyword,
        platform_id=platform_id,
        title=title,
        link=link,
        task_id=task_id,
        is_external=is_external,
    )
    rows = db.execute(stmt).all()
    hit_rows = list(db.execute(select(SensitiveHit)).scalars())
    hit_map: dict[int, list[SensitiveHit]] = {}
    for hit in hit_rows:
        if hit.result_id:
            hit_map.setdefault(hit.result_id, []).append(hit)

    wb = Workbook()
    ws = wb.active
    ws.title = "检索结果"
    headers = [
        "关键词", "平台", "页码", "结果标题", "结果链接", "搜索时间",
        "是否外部链接", "敏感命中级别", "敏感命中分类", "命中文本(脱敏)",
    ]
    ws.append(headers)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F6F3D")
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
    widths = [18, 12, 6, 40, 50, 20, 12, 14, 16, 30]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width

    for result, platform_name in rows:
        hits = hit_map.get(result.id, [])
        level_order = {"high": 0, "medium": 1, "low": 2}
        hits = sorted(hits, key=lambda h: level_order.get(h.level, 3))
        if hits:
            for idx, hit in enumerate(hits):
                ws.append(
                    [
                        result.keyword if idx == 0 else "",
                        platform_name if idx == 0 else "",
                        result.page if idx == 0 else "",
                        result.title if idx == 0 else "",
                        result.url if idx == 0 else "",
                        result.searched_at.strftime("%Y-%m-%d %H:%M:%S")
                        if idx == 0 and result.searched_at
                        else "",
                        {True: "是", False: "否", None: ""}[result.is_external]
                        if idx == 0
                        else "",
                        {"high": "高", "medium": "中", "low": "低"}.get(
                            hit.level, hit.level
                        ),
                        hit.category,
                        hit.matched_text,
                    ]
                )
        else:
            ws.append(
                [
                    result.keyword,
                    platform_name,
                    result.page,
                    result.title,
                    result.url,
                    result.searched_at.strftime("%Y-%m-%d %H:%M:%S")
                    if result.searched_at
                    else "",
                    {True: "是", False: "否", None: ""}[result.is_external],
                    "",
                    "",
                    "",
                ]
            )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"sdscan_results_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
